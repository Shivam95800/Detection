import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_memo_to_docx():
    base_dir = r"C:\Users\Asus\.gemini\antigravity-ide\scratch\kestrel-warranty"
    md_path = os.path.join(base_dir, "MEMO_TO_RITU.md")
    docx_path = os.path.join(base_dir, "MEMO_TO_RITU.docx")
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = Document()
    
    # Page setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    title = doc.add_heading("MEMORANDUM", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(20)
    title.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped == "# MEMORANDUM":
            continue
            
        if stripped.startswith("**TO:**") or stripped.startswith("**FROM:**") or stripped.startswith("**DATE:**") or stripped.startswith("**SUBJECT:**"):
            p = doc.add_paragraph()
            parts = stripped.split("**", 2)
            if len(parts) >= 3:
                r1 = p.add_run(parts[1])
                r1.bold = True
                r1.font.size = Pt(10)
                r2 = p.add_run(parts[2])
                r2.font.size = Pt(10)
        elif stripped == "---":
            # Add divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            h_text = stripped.replace("### ", "")
            h = doc.add_heading(h_text, level=2)
            h.runs[0].font.size = Pt(12)
            h.runs[0].font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            # Handle basic bolding
            parts = text.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(9.5)
                if i % 2 == 1:
                    run.bold = True
            p.paragraph_format.space_after = Pt(2)
        elif stripped.startswith("1. ") or stripped.startswith("2. ") or stripped.startswith("3. ") or stripped.startswith("4. "):
            p = doc.add_paragraph(style='List Number')
            text = stripped[3:]
            parts = text.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(9.5)
                if i % 2 == 1:
                    run.bold = True
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            parts = stripped.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(9.5)
                if i % 2 == 1:
                    run.bold = True
            p.paragraph_format.space_after = Pt(4)
            
    doc.save(docx_path)
    print(f"Exported clean executive DOCX to: {docx_path}")

if __name__ == "__main__":
    export_memo_to_docx()
